from docx import Document
from docx.shared import Inches
from pathlib import Path
import re

root = Path(__file__).resolve().parents[1]
path = root / 'gTheoryShiny_EPM_manuscript.docx'
doc = Document(path)
anchor = next(p for p in doc.paragraphs if p.text.strip() == 'References')
title = anchor.insert_paragraph_before('Interpretive Walkthrough: Linking Output to Measurement Decisions')
title.style = doc.styles['Heading 1']
paragraphs = [
"A useful way to read a G-study output is to move from the intended decision to the error definition, rather than beginning with the largest number in the variance table. Suppose the assessment is used to rank persons. The relevant question is which sources change relative ordering. Person-by-item and person-by-rater interactions are therefore central relative-error terms, while a main item or rater effect may be less important for ranking if it affects all persons similarly. If the assessment is used to classify persons against a standard, the absolute-error coefficient is more relevant because systematic differences among items or raters can move an observed score relative to the standard. gTheoryShiny places both coefficients and both error summaries in the result panel to make this distinction difficult to overlook.",
"The worked example illustrates why this interpretation matters. The person component is about 1.04, but the person-by-item component is about .26 and the person-by-rater component is about .11. Adding items reduces the first interaction by an item denominator; adding raters reduces the second by a rater denominator. The item and rater main effects enter the absolute error as well. Consequently, a design that increases items can improve G more efficiently than Phi when main facet variation is substantial. The D-study table allows the analyst to see this difference directly instead of inferring it from a single reported reliability coefficient.",
"A second interpretive question concerns replication. A random facet term is not supported simply because a column has a name. The application checks whether observations are replicated sufficiently to estimate the requested grouping term. If each person is observed only once per rater, a person-by-rater term may be estimable; if a supposed facet is perfectly confounded with the person or another facet, it may not be. The software reports omitted non-estimable terms and retains the selected data for inspection. The analyst should then decide whether the term was conceptually necessary, whether the design should be changed, or whether the available data cannot answer the intended question.",
"A third question concerns the difference between a warning and a failure. A convergence warning does not automatically invalidate every output, but it identifies a model that requires examination. A singular fit often indicates that one or more variance components are estimated at or near a boundary. In a teaching demonstration, it can be useful to show the warning and compare a simplified model. In a high-stakes assessment, it should prompt additional replication, design simplification, or a sensitivity analysis before a coefficient is used to justify a decision. The application reports warnings in plain language but does not suppress them to make the interface appear cleaner.",
"The same discipline applies to bootstrap intervals. A narrow interval can reflect a stable design, but it can also reflect an estimator that repeatedly returns similar boundary values. The user should inspect the distribution of bootstrap draws, the proportion of failed refits, and whether the fitted model is singular. gTheoryShiny supplies the draw-level file so that analysts can make this inspection in R or another tool. A manuscript should report the bootstrap type, number of draws, resampling unit, percentile or other interval rule, failed draws, and software version. These details are necessary because an interval without its generation procedure is difficult to reproduce.",
"For multivariate analyses, weights should be treated as part of the score interpretation. Equal weights are not automatically appropriate when subtests have different scales, score ranges, or substantive importance. The application asks the user to enter weights explicitly and checks that the number of weights matches the fixed outcomes. The covariance selection is also explicit. An unstructured matrix can represent shared error or facet covariance, but it consumes more information than a diagonal matrix. The walkthrough therefore recommends comparing structures, checking convergence, and explaining why the selected covariance model corresponds to the assessment purpose.",
"The AI Assistant can support this interpretive process by translating displayed results into questions, but it cannot resolve design ambiguity. A useful prompt asks the assistant to identify the largest nonperson component and suggest which D-study condition would reduce it. An inappropriate prompt asks the assistant to declare that an assessment is valid or to choose a cut score without context. The privacy boundary also matters: even aggregate context can be sensitive when variable names or small sample sizes identify a project. Users should configure the optional service only when their institutional data policy permits it and should review generated answers before sharing them.",
"A final interpretive question is whether the software output is sufficient for a published claim. It is not. A coefficient is evidence about score consistency under a specified design and decision rule. Validity concerns the interpretation and use of scores, including construct representation, response processes, relations to other variables, and consequences. gTheoryShiny can make the reliability part more transparent, but the surrounding validity argument remains the responsibility of the researchers. The application is most valuable when its tables, warnings, and downloaded artifacts are incorporated into that broader argument rather than presented as an isolated software result."
]
for text in paragraphs:
    p = anchor.insert_paragraph_before(text)
    p.paragraph_format.first_line_indent = Inches(.5)
    p.paragraph_format.line_spacing = 2
doc.save(path)
blind = Document(path)
for p in blind.paragraphs:
    if p.text.startswith(('Zhehan Jiang, Jihong Zhang, and Tianpeng Zheng','Institute of Medical Education','Corresponding authors:','Jihong Zhang is co-first','Author Note.')):
        p.text = '[Author information removed for blinded review]'
blind.core_properties.author = 'Anonymous'
blind.save(root / 'gTheoryShiny_EPM_manuscript_blinded.docx')
texts = [p.text for p in doc.paragraphs]
i = texts.index('References')
print('body_words', len(re.findall(r"\b[\w’\'-]+\b", ' '.join(texts[:i]))), 'references', len(texts)-i-1)
